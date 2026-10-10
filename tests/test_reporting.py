"""Tests for the HTML and Excel reports."""
from io import BytesIO

import numpy as np
import pandas as pd
from openpyxl import load_workbook

from src.ai.insights import generate_insights
from src.analysis.pipeline import run_analysis
from src.cleaning.cleaner import clean_dataframe
from src.profiling.type_detector import profile_columns
from src.reporting.excel_report import build_excel_report
from src.reporting.html_report import build_html_report


def make_inputs(regions=("North", "South", "East")):
    rng = np.random.default_rng(3)
    n = 400
    df = pd.DataFrame(
        {
            "order_date": pd.date_range("2024-01-01", periods=n, freq="D"),
            "revenue": 100 + 0.25 * np.arange(n) + rng.normal(0, 3, n),
            "units": rng.integers(1, 20, n),
            "region": rng.choice(list(regions), n),
        }
    )
    df.loc[[20, 120, 220], "revenue"] = 1000.0
    cleaning = clean_dataframe(df)
    bundle = run_analysis(cleaning)
    insights = generate_insights(bundle, cleaning, "x.csv", use_llm=False)
    return df, cleaning, bundle, insights


def html_for(title="Test report", scope="x.csv", mode="none") -> str:
    df, cleaning, bundle, insights = make_inputs()
    return build_html_report(
        title=title, scope_label=scope, cleaning=cleaning, full_cleaning=cleaning,
        raw_profiles=profile_columns(df), bundle=bundle, insights=insights,
        filters_applied=[], chart_mode=mode,
    )


def test_html_report_contains_all_sections():
    html = html_for()
    for heading in [
        "Executive summary", "Key metrics", "Important findings", "Insights and recommendations",
        "Trends", "Unusual values", "Data and cleaning", "How this report was made",
    ]:
        assert f"<h2>{heading}</h2>" in html
    assert "<title>Test report</title>" in html


def test_html_report_escapes_text_from_data():
    html = html_for(title="<script>alert(1)</script>", scope="<b>x</b>")
    assert "<script>" not in html
    assert "&lt;script&gt;alert(1)&lt;/script&gt;" in html
    assert "<b>x</b>" not in html


def test_html_report_with_charts_embeds_plotly():
    html = html_for(mode="inline")
    assert "Plotly.newPlot" in html
    assert len(html) > 1_000_000  # plotly.js is embedded once, so the file works offline


def test_excel_report_has_sheets_and_blocks_formulas():
    _, cleaning, bundle, insights = make_inputs(regions=("North", "South", "=SUM(1,2)"))
    data = build_excel_report(
        title="Test", scope_label="x.csv", cleaning=cleaning, full_cleaning=cleaning,
        bundle=bundle, insights=insights, filters_applied=[],
    )
    workbook = load_workbook(BytesIO(data))
    for name in ["Summary", "KPIs", "Findings", "Insights", "Cleaned data"]:
        assert name in workbook.sheetnames

    sheet = workbook["Cleaned data"]
    values = [c for row in sheet.iter_rows() for c in row]
    assert all(c.data_type != "f" for c in values)  # nothing is stored as a formula
    assert any(isinstance(c.value, str) and c.value.startswith("'=SUM") for c in values)